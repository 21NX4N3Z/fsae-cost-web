"""
ตรวจว่าสูตรทุกเซลล์อ้างถึงแถวที่ถูกต้อง

openpyxl ไม่ปรับสูตรให้เมื่อ insert_rows ทำให้สูตรชี้ไปแถวอื่นเงียบ ๆ
สคริปต์นี้จับกรณีนั้น: ทุกสูตรที่อยู่ในช่วงข้อมูลของหมวด ต้องอ้างแถว
ที่อยู่ในช่วงเดียวกัน (หรือแถวยอดรวมของหมวด) ไม่งั้นคือสูตรเลื่อนไปแล้ว

รัน: python scripts/verify_formulas.py
"""
import re
import openpyxl
from openpyxl.utils import get_column_letter as gl

PATH = r"D:/hermes-workspace/fsae-cost-web/data/BP18 FR Cost.xlsx"

# หัวตารางของแต่ละหมวด -> คอลัมน์ที่เก็บยอดรวม
SECTIONS = {
    "Material": ("I", "AA"),
    "Process": ("I", "S"),
    "Fastener": ("I", "U"),
    "Tooling": ("I", "S"),
}
DETAILS_HEADER = {"L": "Item Order", "N": "Part", "R": "Part Cost",
                  "W": "Quantity", "Y": "Subtotal"}


def col_idx(letter):
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch) - 64)
    return n


def find_sections(ws):
    """คืนรายการ (ชื่อ, แถวหัว, แถวยอดรวม) ของทุกหมวดในแท็บ"""
    out = []
    for name, (ucol, tcol) in SECTIONS.items():
        for r in range(1, ws.max_row + 1):
            if str(ws.cell(r, col_idx(ucol)).value).strip() == "Unit Cost" \
               and str(ws.cell(r, col_idx(tcol)).value).strip().lower() == "sub total":
                # แถวยอดรวมคือแถวถัดไปที่มีป้าย Sub Total
                sub = None
                for rr in range(r + 1, min(r + 80, ws.max_row + 1)):
                    v = str(ws.cell(rr, col_idx(tcol) - 2).value).strip().lower()
                    if v == "sub total":
                        sub = rr
                        break
                    if str(ws.cell(rr, col_idx(ucol)).value).strip() == "Unit Cost":
                        break
                out.append((name, r, sub, ucol, tcol))
    return out


def find_details(ws):
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, col_idx("L")).value).strip() == "Item Order" \
           and str(ws.cell(r, col_idx("N")).value).strip() == "Part":
            sub = None
            for rr in range(r + 1, ws.max_row + 1):
                if str(ws.cell(rr, col_idx("W")).value).strip().lower() == "subtotal":
                    sub = rr
                    break
                if str(ws.cell(rr, col_idx("L")).value).strip() == "Item Order":
                    break
            return (r, sub)
    return (None, None)


def main():
    wb = openpyxl.load_workbook(PATH)
    bad = []
    checked = 0
    dangling = 0
    names = set(wb.sheetnames)

    for name in wb.sheetnames:
        ws = wb[name]
        secs = find_sections(ws)
        dh, dsub = find_details(ws)

        # เซลล์ที่ต้องอยู่ในช่วงใด -> (ช่วงเริ่ม, ช่วงจบ, ชื่อช่วง)
        zones = []
        if dh:
            zones.append((dh + 1, (dsub - 1) if dsub else ws.max_row, "Details"))
        for sname, hr, sub, ucol, tcol in secs:
            end = (sub - 1) if sub else hr + 30
            zones.append((hr + 1, end, sname))

        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or not v.startswith("="):
                    continue
                checked += 1

                # อ้างแท็บที่ไม่มีอยู่จริง
                for t in re.findall(r"'([^']+)'!", v):
                    if t not in names:
                        dangling += 1
                        bad.append((name, cell.coordinate, "อ้างแท็บที่ไม่มี: " + t, v[:70]))

                # สูตรในชีตนี้เอง -> ต้องอยู่ในช่วงที่ตัวเองอยู่
                if "!" in v:
                    continue
                owner = None
                for z0, z1, zname in zones:
                    if z0 <= cell.row <= z1:
                        owner = (z0, z1, zname)
                        break
                if not owner:
                    continue
                z0, z1, zname = owner
                for m in re.finditer(r"([A-Z]{1,2})(\d+)", v):
                    ref = int(m.group(2))
                    if not (z0 <= ref <= z1 + 1):
                        bad.append((name, cell.coordinate,
                                    "%s: อ้างแถว %d นอกช่วง %d-%d" % (zname, ref, z0, z1), v[:70]))

    print("ตรวจสูตร %d สูตร ใน %d แท็บ" % (checked, len(wb.sheetnames)))
    print("อ้างแท็บที่ไม่มีอยู่จริง: %d" % dangling)
    print("สูตรที่อ้างแถวผิดช่วง: %d" % (len(bad) - dangling))
    if bad:
        print("")
        for b in bad[:25]:
            print("  !! %-14s %-6s %s" % (b[0], b[1], b[2]))
            print("     %s" % b[3])
        if len(bad) > 25:
            print("  ... อีก %d รายการ" % (len(bad) - 25))
    return 0 if not bad else 1


if __name__ == "__main__":
    raise SystemExit(main())
