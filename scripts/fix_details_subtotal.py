"""
แก้ยอดรวม Details ของแท็บแม่ที่เป็นตัวเลขตาย ให้เป็นสูตร SUM

ต้นทาง (BP16b) บางแท็บใส่ยอดรวมเป็นตัวเลข เช่น FR 00300-AA = 106.98
ถ้าทีมแก้ต้นทุนชิ้นส่วน ยอดของแม่จะไม่เปลี่ยนตาม
แท็บที่ไม่มีหัวตาราง Details เลย (เช่น ANTI INTRUSION PLATE) ข้ามไป
"""
import re
import openpyxl

PATH = r"D:/hermes-workspace/fsae-cost-web/data/BP18 FR Cost.xlsx"
DET_HEADER_COL = 12   # L
DET_PART_COL = 14     # N
DET_TOTAL_LABEL_COL = 23   # W = 'Subtotal'
DET_VALUE_COL = 25    # Y


def fix():
    wb = openpyxl.load_workbook(PATH)
    fixed = []
    skipped = []

    for name in wb.sheetnames:
        ws = wb[name]
        # หาแถวหัว Details
        header = None
        for r in range(1, ws.max_row + 1):
            if str(ws.cell(r, DET_HEADER_COL).value).strip() == "Item Order" \
               and str(ws.cell(r, DET_PART_COL).value).strip() == "Part":
                header = r
                break
        if header is None:
            skipped.append(name)
            continue

        # หาแถวยอดรวม
        total = None
        for r in range(header + 1, ws.max_row + 1):
            if str(ws.cell(r, DET_TOTAL_LABEL_COL).value).strip().lower() == "subtotal":
                total = r
                break
            if str(ws.cell(r, DET_HEADER_COL).value).strip() == "Item Order":
                break
        if total is None:
            skipped.append(name)
            continue

        cur = ws.cell(total, DET_VALUE_COL).value
        if isinstance(cur, str) and cur.strip().startswith("="):
            continue

        first, last = header + 1, total - 1
        if last < first:
            skipped.append(name)
            continue
        ws.cell(total, DET_VALUE_COL).value = "=SUM(Y%d:Y%d)" % (first, last)
        fixed.append("%s  Y%d  %s -> =SUM(Y%d:Y%d)" % (name, total, cur, first, last))

    wb.save(PATH)
    print("แก้ยอดรวม Details: %d แท็บ" % len(fixed))
    for f in fixed:
        print("  " + f)
    print("ไม่มี Details (ข้าม): %d แท็บ" % len(skipped))
    return 0


if __name__ == "__main__":
    raise SystemExit(fix())
